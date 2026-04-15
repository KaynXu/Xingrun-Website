from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def build_roster(input_path: Path) -> dict[str, dict[str, list[str]]]:
    workbook = load_workbook(input_path, data_only=True)
    roster: dict[str, dict[str, list[str]]] = defaultdict(dict)

    for sheet in workbook.worksheets:
        teacher = normalize_text(sheet["C2"].value)
        class_name = normalize_text(sheet["G2"].value) or sheet.title.strip()

        if not teacher:
            continue

        students: list[str] = []
        seen: set[str] = set()

        for row in range(5, sheet.max_row + 1):
            name = normalize_text(sheet[f"B{row}"].value)
            if not name or name in seen:
                continue
            seen.add(name)
            students.append(name)

        roster[teacher][class_name] = students

    return {
        teacher: {
            class_name: classes[class_name]
            for class_name in sorted(classes.keys(), key=lambda item: item)
        }
        for teacher, classes in sorted(roster.items(), key=lambda item: item[0])
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate offline classroom roster JS data from an xlsx file.")
    parser.add_argument("input", type=Path, help="Path to the source xlsx file.")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("classroom-roster-data.js"),
        help="Path to the output JS file.",
    )
    args = parser.parse_args()

    roster = build_roster(args.input)
    total_classes = sum(len(classes) for classes in roster.values())
    total_students = sum(len(students) for classes in roster.values() for students in classes.values())

    output = (
        "window.CLASSROOM_ROSTER_META = "
        + json.dumps(
            {
                "source_file": str(args.input),
                "teacher_count": len(roster),
                "class_count": total_classes,
                "student_count": total_students,
            },
            ensure_ascii=False,
            indent=2,
        )
        + ";\n\nwindow.CLASSROOM_ROSTER_DATA = "
        + json.dumps(roster, ensure_ascii=False, indent=2)
        + ";\n"
    )

    args.output.write_text(output, encoding="utf-8")
    print(f"Generated {args.output} with {len(roster)} teachers, {total_classes} classes, {total_students} students.")


if __name__ == "__main__":
    main()
